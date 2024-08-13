import openpyxl
from pywinauto.application import Application
import os
import time

# Format the column of phone numbers as text
def format_phone_column_as_text(excel_file, sheet_name, column_letter):
    try:
        print("Starting to format the phone column as text...")
        wb = openpyxl.load_workbook(excel_file)
        ws = wb[sheet_name]

        # Apply the text format to each cell in the column
        for cell in ws[column_letter]:
            cell.number_format = '@'

        # Save the sheet with the new format
        wb.save(excel_file)
        print(f"Column {column_letter} formatted as text successfully.")
    except Exception as e:
        print(f"Couldn't format the column of phone numbers: {e}")

# Read names n numbers of contacts from the Excel
def get_contacts_from_excel(excel_file, sheet_name):
    contacts = []
    try:
        print("Reading contacts from Excel...")
        wb = openpyxl.load_workbook(excel_file)
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):  # Assuming data starts from row 2
            name, number = row
            if name and number:
                contacts.append((name, number))
        print(f"Contacts read from Excel: {contacts}")
        return contacts
    except Exception as e:
        print(f"Couldn't read the contacts from Excel: {e}")
        return []

# Update STATUS in Excel
def update_status_in_excel(excel_file, sheet_name, row, status):
    try:
        print(f"Updating status in Excel for row {row}...")
        wb = openpyxl.load_workbook(excel_file)
        ws = wb[sheet_name]
        status_column = 'C'  # Status Column
        ws[f'{status_column}{row}'] = status
        wb.save(excel_file)
        print(f"Status updated on row {row}: {status}")
    except Exception as e:
        print(f"Couldn't update status on row {row}: {e}")

# Send a PDF via pywinauto to a specific contact
def send_pdf(app, contact_name, pdf_path):
    try:
        print(f"Attempting to send PDF to '{contact_name}'...")
        # Ensure the WhatsApp window is in the foreground
        app.top_window().set_focus()

        # Find and interact with search box
        search_box = app.top_window().child_window(auto_id="SearchBox", control_type="Edit")
        search_box.click_input()
        search_box.type_keys(contact_name, with_spaces=True)

        time.sleep(2)  # Wait for search results to appear

        # Select the contact
        contact = app.top_window().child_window(title=contact_name, control_type="Text")
        contact.click_input()

        # Attach file
        attach_button = app.top_window().child_window(title="Attach", control_type="Button")
        attach_button.click_input()

        # Open File Dialog and send PDF
        app.Dialog.child_window(title="Document", control_type="Text").click_input()
        time.sleep(1)
        app.Dialog.child_window(auto_id="1148", control_type="Edit").type_keys(pdf_path)
        app.Dialog.child_window(auto_id="1", control_type="Button").click_input()

        time.sleep(2)  # Wait for the file to upload

        # Send the PDF
        send_button = app.top_window().child_window(title="Send", control_type="Button")
        send_button.click_input()

        print(f"PDF '{pdf_path}' sent to '{contact_name}' successfully!")
        return True

    except Exception as e:
        print(f"Couldn't send the file '{pdf_path}' to '{contact_name}': {e}")
        return False

def process_and_send_pdfs(excel_file, sheet_name, pdf_directory):
    contacts = get_contacts_from_excel(excel_file, sheet_name)
    if not contacts:
        print("No contacts found to send the PDF files.")
        return

    try:
        print("Launching WhatsApp Desktop application...")
        # Launch the WhatsApp Desktop app
        app = Application().start("shell:AppsFolder\\5319275A.WhatsAppDesktop_cv1g1gvanyjgm!App")
        time.sleep(10)  # Wait for WhatsApp to load

        for index, (name, _) in enumerate(contacts, start=2):  # Assuming data starts from row 2
            print(f"Processing contact '{name}' (row {index})...")
            pdf_name = f"{name}.pdf"
            pdf_path = os.path.join(pdf_directory, pdf_name)

            if os.path.isfile(pdf_path):
                print(f"PDF found for '{name}' at '{pdf_path}'.")
                success = send_pdf(app, name, pdf_path)
                status = "Successfully sent" if success else "Couldn't send the file"
            else:
                print(f"PDF file not found for '{name}' at '{pdf_path}'.")
                status = f"File not found: {pdf_path}"

            update_status_in_excel(excel_file, sheet_name, index, status)

    except Exception as e:
        print(f"General issue encountered: {e}")

# Main script execution
excel_file = 'C:\\Users\\Carolina\\Desktop\\Automate Whatsapp-Pdf- Excel\\Excel_Python_Project.xlsx'
sheet_name = 'Sheet1'
column_letter = 'B'  # 2nd column has mobile numbers
pdf_directory = r'C:\Users\Carolina\Desktop\Automate Whatsapp-Pdf- Excel\pdfs'

print("Starting the PDF sending process...")
format_phone_column_as_text(excel_file, sheet_name, column_letter)
process_and_send_pdfs(excel_file, sheet_name, pdf_directory)
print("PDF sending process completed.")
